#!/usr/bin/env python3
"""Generator arkusza QA dla EdSharp 5.0.1 (wlasna galaz na bazie 5.0 beta).

Format ustalony z Michalem (jak w 4.0.x): zakladki Retest / Instrukcja /
Podsumowanie / Dane testowe; komunikaty aplikacji po angielsku, reszta po polsku.
Uruchom: ~/.hermes/hyperspace_venv/bin/python build_qa_sheet_5.0.1.py
"""
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation

WERSJA = "5.0.1"
BUILD = WERSJA

# [ID, Obszar, Priorytet, Status implementacji, Zrodlo, Kroki,
#  Oczekiwany rezultat, None(wynik), None(uwagi), Build]
CASES = [
    # --- Zadania z "Do zrobienia 2026-08-04.md" ---
    ["HTML-001", "Usuniecie konwersji do HTML (Ctrl+H)", "P1",
     "Zrobione: usunieto komende HTML Format - skrot Ctrl+H, pozycje w menu Misc, handler i helper Markdown2Html.",
     "Do zrobienia 2026-08-04.md (zad. 'Zmiana konwersji na HTML')",
     "1. Otworz dowolny plik .md. 2. Wcisnij Ctrl+H. 3. Otworz menu Misc i przejrzyj pozycje.",
     "Ctrl+H NIC nie robi (brak konwersji, plik zrodlowy nietkniety). W menu Misc NIE ma pozycji 'HTML Format'. Konwersja do HTML dostepna tylko przez Zapisz jako.",
     None, None, BUILD],
    ["HEAD-001", "Wstawianie naglowka 6 (Ctrl+Shift+6)", "P1",
     "Zrobione: przeniesiono kolidujaca komende 5.0 'Prior Baseline' z Ctrl+Shift+D6 na Ctrl+Alt+D6, dzieki czemu Ctrl+Shift+6 dziala.",
     "Do zrobienia 2026-08-04.md (zad. 'Wstawianie naglowkow')",
     "1. W pliku .md ustaw kursor w linii z tekstem. 2. Wcisnij Ctrl+Shift+6.",
     "Linia zamienia sie na naglowek poziomu 6 (######  przed tekstem), NVDA oglasza 'Heading 6'. Wczesniej klawisz byl polykany.",
     None, None, BUILD],
    ["HEAD-002", "Wstawianie naglowkow 1-5 (Ctrl+Shift+1..5)", "P1",
     "Bez zmian - dzialalo; regresja po przeniesieniu skrotow.",
     "Do zrobienia 2026-08-04.md",
     "1. W pliku .md w linii z tekstem wcisnij kolejno Ctrl+Shift+1, +2, +3, +4, +5.",
     "Kazdy skrot wstawia odpowiednia liczbe # i oglasza 'Heading N'. Zaden nie jest polykany.",
     None, None, BUILD],
    ["HEAD-003", "Czyszczenie formatowania (Ctrl+Shift+0)", "P1",
     "Zrobione: przeniesiono kolidujaca komende 5.0 'Go to Special Folder' z Ctrl+Shift+D0 na Ctrl+Alt+D0.",
     "Do zrobienia 2026-08-04.md (naglowki i podobne markdownowe)",
     "1. Zaznacz sformatowany tekst (np. naglowek/pogrubienie). 2. Wcisnij Ctrl+Shift+0.",
     "Formatowanie zostaje usuniete, NVDA oglasza 'Clear formatting'. Wczesniej klawisz byl polykany.",
     None, None, BUILD],
    ["CONFLICT-001", "Przeniesione komendy 5.0 nadal dzialaja", "P2",
     "Zrobione: Prior Baseline -> Ctrl+Alt+D6, Go to Special Folder -> Ctrl+Alt+D0; obie nadal w menu.",
     "Skutek uboczny naprawy naglowkow",
     "1. Wcisnij Ctrl+Alt+6 (Prior Baseline). 2. Wcisnij Ctrl+Alt+0 (Go to Special Folder). 3. Sprawdz te pozycje w menu Navigate/Misc.",
     "Prior Baseline i Go to Special Folder dzialaja na nowych skrotach i sa poprawnie opisane w menu.",
     None, None, BUILD],
    ["PREVIEW-001", "Podglad z synchronizacja (Escape)", "P1",
     "Bez zmian - tryb synchronizowany zachowany.",
     "Do zrobienia 2026-08-04.md (kontekst zad. Shift-Esc)",
     "1. W pliku .md ustaw kursor w konkretnym miejscu. 2. Wcisnij Escape (wejscie w podglad). 3. Poruszaj sie po podgladzie strzalkami. 4. Wcisnij Escape (powrot do edycji).",
     "Escape przelacza podglad/edycje. W tym trybie kursor edycji PODAZA za kursorem w podgladzie (synchronizacja jak dotad). NVDA oglasza 'Preview' (wejscie) / 'Editing' (wyjscie) - komunikat ma zgadzac sie z faktycznym stanem.",
     None, None, BUILD],
    ["PREVIEW-002", "Podglad ODCZEPIONY bez synchronizacji (Shift+Escape)", "P1",
     "Zrobione: nowy tryb detached - kursor w podgladzie NIE rusza kursora edycji; po wyjsciu kursor wraca do miejsca sprzed wejscia.",
     "Do zrobienia 2026-08-04.md (zad. 'Podglad bez synchronizacji fokusa')",
     "1. W pliku .md ustaw kursor w KONKRETNYM miejscu (zapamietaj je). 2. Wcisnij Shift+Escape. 3. Przejdz w INNE miejsce podgladu. 4. Wcisnij ponownie Shift+Escape (powrot).",
     "NVDA oglasza 'Preview detached'. Ruch w podgladzie NIE przenosi kursora edycji. Po Shift+Escape kursor edycji jest DOKLADNIE tam, gdzie przed wejsciem w podglad.",
     None, None, BUILD],
    ["PREVIEW-003", "Przelaczanie trybow w miejscu (bezpiecznik)", "P2",
     "Zrobione: bedac w podgladzie, drugi skrot przelacza tryb sync<->detached zamiast wychodzic.",
     "Do zrobienia 2026-08-04.md (uwaga o zabezpieczeniu ESC/Shift-ESC)",
     "1. Wejdz w podglad przez Escape (synchronizowany). 2. Bedac w podgladzie wcisnij Shift+Escape.",
     "Tryb przelacza sie na ODCZEPIONY w miejscu (NVDA: 'Detached'), bez wychodzenia z podgladu. Analogicznie z podgladu odczepionego Escape -> 'Synced'.",
     None, None, BUILD],
    ["FOCUS-001", "Stabilnosc fokusa po Alt+Tab", "P1",
     "Zrobione: handler MdiFrame.Activated wymusza fokus na kontrolce edycyjnej po odzyskaniu okna.",
     "Do zrobienia 2026-08-04.md (zad. 'Widocznosc, stabilnosc fokusa ... Alt-Tab')",
     "1. Otworz plik w EdSharp. 2. Przelacz sie Alt+Tab na inne okno i z powrotem na EdSharp (kilka razy, rozne typy plikow). 3. Od razu zacznij czytac/pisac.",
     "Po powrocie fokus i mowa NVDA sa OD RAZU w tekscie (bez zamrozenia, bez potrzeby ponownego wejscia czy podwojnego Escape). UWAGA: objaw moze byc srodowiskowy - odnotuj czy wystepuje.",
     None, None, BUILD],
    ["FOCUS-002", "Fokus po Alt+Tab w trybie podgladu", "P2",
     "Zrobione: handler kieruje fokus na widok podgladu gdy aktywny.",
     "Skutek uboczny FOCUS-001",
     "1. Wejdz w podglad (Escape) w pliku .md. 2. Alt+Tab na inne okno i z powrotem.",
     "Po powrocie fokus jest w widoku podgladu (mowa dziala od razu), tryb podgladu zachowany.",
     None, None, BUILD],
    # --- Regresja skonsolidowanych funkcji (nasza galaz na 5.0) ---
    ["REG-COPY-001", "Rich copy Markdown (Ctrl+Shift+C)", "P2",
     "Skonsolidowane z galezi feature/markdown-editing.",
     "Regresja po przejsciu na baze 5.0",
     "1. W .md zaznacz liste punktowana. 2. Ctrl+Shift+C. 3. Wklej do Worda.",
     "Wkleja sie jako prawdziwa lista Worda; NVDA oglasza 'List copied'. Linki jako hiperlacza ('Link copied').",
     None, None, BUILD],
    ["REG-SECT-001", "Przenoszenie sekcji (Ctrl+Alt+strzalka)", "P2",
     "Skonsolidowane z feature/markdown-section-move.",
     "Regresja po przejsciu na baze 5.0",
     "1. W .md z kilkoma naglowkami ustaw kursor w sekcji. 2. Ctrl+Alt+strzalka w gore / w dol.",
     "Cala sekcja przenosi sie ponad/pod sasiednia; NVDA oglasza 'Above/Below ...', 'Top!'/'Bottom!' na krancach.",
     None, None, BUILD],
    ["REG-LIST-001", "Lista elementow podgladu (F7)", "P2",
     "Skonsolidowane z feature/markdown-review-mode.",
     "Regresja po przejsciu na baze 5.0",
     "1. Wejdz w podglad .md (Escape). 2. Wcisnij F7. 3. Wybierz element.",
     "Otwiera sie lista elementow (naglowki/linki/listy/tabele); wybor skacze do elementu w podgladzie.",
     None, None, BUILD],
    ["REG-NAV-001", "Nawigacja w podgladzie (h/l/i/k/t)", "P2",
     "Skonsolidowane z feature/markdown-review-mode.",
     "Regresja po przejsciu na baze 5.0",
     "1. W podgladzie .md wcisnij h (naglowek), l (lista), i (element listy), k (link), t (tabela); Shift+klawisz = wstecz.",
     "Kursor skacze do kolejnego/poprzedniego elementu danego typu, NVDA oglasza element.",
     None, None, BUILD],
    ["REG-OPEN-001", "Open With z opcja 'zawsze'", "P2",
     "Skonsolidowane z feature/open-with-dialog.",
     "Regresja po przejsciu na baze 5.0",
     "1. Na liscie Recent/Favorites albo z menu wywolaj 'Open With' na pliku.",
     "Pojawia sie nowoczesny dialog 'Otworz za pomoca' z opcja 'Zawsze uzywaj tej aplikacji'.",
     None, None, BUILD],
    ["REG-FILELIST-001", "Akcje na listach Recent/Favorites", "P2",
     "Skonsolidowane z feature/file-list-actions-v5 (reimplementacja na LbcDialog 5.0).",
     "Regresja po przejsciu na baze 5.0",
     "1. Otworz liste Recent lub Favorites. 2. Na wpisie: strzalka w prawo (Open With), w lewo (odczyt sciezki), Ctrl+Enter (pokaz w Eksploratorze), Ctrl+C (kopiuj sciezke), Delete (usun z listy), Shift+Delete (usun z dysku).",
     "Kazda akcja dziala; Shift+Delete pyta o potwierdzenie i NIE kasuje gdy zamkniecie okna anulowano ('Delete canceled'), inaczej 'Deleted from disk'.",
     None, None, BUILD],
    ["NVDA-SPELL-001", "Dodatek NVDA: bledy pisowni", "P1",
     "Nasza wtyczka edsharpngSpellcheck przeniesiona do 5.0; appModule 'edsharpng' pasuje do procesu EdSharpNG.exe.",
     "Nowa funkcja / przeniesienie",
     "1. W instalatorze na ekranie koncowym zaznacz 'Install NVDA spelling-errors add-on for EdSharpNG'. 2. Zrestartuj NVDA. 3. Otworz EdSharpNG, wpisz tekst z bledem (np. 'kot pisez'). 4. Przejdz po tekscie strzalkami/czytaniem.",
     "NVDA sygnalizuje blednie napisane slowa (jak w Wordzie), w jezyku pisowni Windows. W logu NVDA (poziom Debug) wpis 'EdSharpNG spellcheck:'.",
     None, None, BUILD],
    ["NVDA-PASS-001", "Dodatek NVDA: pass-through klawiszy (Jamal)", "P2",
     "Pass-through Jamala dolaczony do instalatora; oddaje klawisze Matrix/list do aplikacji.",
     "Regresja / dodatek dolaczony",
     "1. W instalatorze zaznacz 'Install NVDA add-on'. 2. Zrestartuj NVDA. 3. W EdSharpNG uzyj Ctrl+Alt+strzalki/Home/End (nawigacja Matrix) oraz Alt+strzalka gora/dol (listy).",
     "NVDA NIE przechwytuje tych skrotow swoimi komendami; trafiaja do EdSharpNG. Dodatek sam nie mowi (EdSharpNG mowi przez UIA). Obie wtyczki moga byc zainstalowane naraz.",
     None, None, BUILD],
    ["BUILD-002", "Zmiana nazwy pliku wykonywalnego na EdSharpNG.exe", "P1",
     "Exe zmieniony EdSharp.exe -> EdSharpNG.exe (dopasowanie do appModule NVDA). Config EdSharpNG.exe.config obok binarki.",
     "Zmiana nazwy exe",
     "1. Po instalacji sprawdz w katalogu programu obecnosc EdSharpNG.exe i EdSharpNG.exe.config (brak starego EdSharp.exe). 2. Uruchom z menu Start i ze skrotu na pulpicie. 3. W Menedzerze zadan sprawdz nazwe procesu.",
     "Program uruchamia sie z obu skrotow; proces nazywa sie EdSharpNG.exe; startowe strojenie dziala (brak zawieszenia startu na sprawdzaniu certyfikatow). Deinstalacja usuwa EdSharpNG.exe i klucz App Paths.",
     None, None, BUILD],
    ["BUILD-001", "Uruchomienie i wersja", "P1",
     "Nowa galaz edsharp5-consolidated (5.0 beta + nasze 5 funkcji + 4 zadania), rebrand EdSharpNG.",
     "Sanity",
     "1. Zainstaluj EdSharpNG_Setup.exe. 2. Uruchom EdSharpNG (Alt+Ctrl+E). 3. Sprawdz tytul okna i Help > About.",
     "Aplikacja startuje na Windows 10+/x64 (i ARM64), NVDA czyta interfejs, brak bledow startowych. Tytul okna i About pokazuja 'EdSharpNG 5.0.1'. Istniejace ustawienia (Recent/Favorites) zachowane.",
     None, None, BUILD],
]

INSTRUKCJA = [
    ["Artefakt", "EdSharpNG_Setup.exe (Inno Setup, x64/ARM64). Baza: EdSharp 5.0 beta + nasze funkcje. Plik wykonywalny: EdSharpNG.exe."],
    ["Wersja", WERSJA + " (galaz edsharp5-consolidated)"],
    ["Cel", "Retest 4 zadan z 'Do zrobienia 2026-08-04.md' + regresja skonsolidowanych funkcji + zmiana nazwy exe na EdSharpNG.exe + 2 dodatki NVDA (pisownia, pass-through)."],
    ["Statusy", "PASS = dziala, FAIL = blad, BLOCKED = brak warunkow, N/A = nie dotyczy."],
    ["Uwaga jezykowa", "Aplikacja po angielsku; arkusz po polsku. Komunikaty (Heading, Preview, Preview detached, Editing, List copied, Top!) zostaja po angielsku."],
    ["Dodatki NVDA", "Instalator oferuje na ekranie koncowym DWA dodatki (domyslnie odznaczone): 1) pisownia (edsharpngSpellcheck - nasza), 2) pass-through klawiszy (Jamala). Po zaznaczeniu NVDA musi byc uruchomione; po instalacji zrestartuj NVDA. Mozna zainstalowac obie naraz."],
    ["Lista statusow", "Kolumna 'Wynik testera': lista PASS/FAIL/BLOCKED/N/A. Nie wpisuj innych wartosci."],
    ["Srodowisko", "Windows 10+ (x64/ARM64) + NVDA. Microsoft Word dla przypadku rich copy. NVDA uruchom przed EdSharp."],
    ["Skroty kluczowe", "Ctrl+Shift+1..6 = naglowki, Ctrl+Shift+0 = clear format, Escape = podglad sync, Shift+Escape = podglad odczepiony, F7 = lista elementow w podgladzie, Ctrl+Alt+strzalki = przenies sekcje."],
    ["Uwaga do FOCUS-001", "Michal zglaszal mozliwe zamrozenie fokusa/NVDA przy Alt+Tab; moze byc srodowiskowe. Odnotuj dokladnie czy i kiedy wystepuje."],
]

DANE = [
    ["Lista punktowana", "- pierwszy element\n- drugi element\n- trzeci element"],
    ["Lista numerowana", "1. Pierwszy\n2. Drugi\n3. Trzeci"],
    ["Polskie znaki", "- zolc i ges\n- cma, zdzblo"],
    ["Naglowki + sekcje", "# Tytul\n\n## Sekcja 1\ntresc 1\n\n## Sekcja 2\ntresc 2\n\n### Podsekcja 2.1\ntresc"],
    ["Link + tabela (podglad/F7)", "[EdSharp](https://github.com/JamalMazrui/EdSharp)\n\n| A | B |\n|---|---|\n| 1 | 2 |"],
]


def build():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = f"Retest {WERSJA}"
    headers = ["ID", "Obszar", "Priorytet", "Status implementacji", "Zrodlo zgloszenia",
               "Kroki retestu", "Oczekiwany rezultat", "Wynik testera", "Uwagi testera", "Build"]
    ws.append(headers)
    for r in CASES:
        ws.append(r)

    hdr_fill = PatternFill("solid", fgColor="DDDDDD")
    thin = Side(style="thin", color="BBBBBB")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    for c in ws[1]:
        c.fill = hdr_fill
        c.font = Font(bold=True)
        c.alignment = Alignment(vertical="center", wrap_text=True)
    for col, w in {"A": 14, "B": 34, "C": 10, "D": 50, "E": 40, "F": 64, "G": 66, "H": 14, "I": 30, "J": 10}.items():
        ws.column_dimensions[col].width = w
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(vertical="top", wrap_text=True)
            c.border = border
    ws.freeze_panes = "A2"
    ws.row_dimensions[1].height = 30
    dv = DataValidation(type="list", formula1='"PASS,FAIL,BLOCKED,N/A"', allow_blank=True)
    dv.add("H2:H200")
    ws.add_data_validation(dv)

    wi = wb.create_sheet("Instrukcja")
    for r in INSTRUKCJA:
        wi.append(r)
    wi.column_dimensions["A"].width = 20
    wi.column_dimensions["B"].width = 110
    for row in wi.iter_rows():
        row[0].font = Font(bold=True)
        for c in row:
            c.alignment = Alignment(vertical="top", wrap_text=True)

    wp = wb.create_sheet("Podsumowanie")
    S = f"'Retest {WERSJA}'"
    for r in [["Metryka", "Wartosc"],
              ["Liczba przypadkow", f"=COUNTA({S}!A2:A200)"],
              ["P1", f'=COUNTIF({S}!C2:C200,"P1")'],
              ["P2", f'=COUNTIF({S}!C2:C200,"P2")'],
              ["PASS", f'=COUNTIF({S}!H2:H200,"PASS")'],
              ["FAIL", f'=COUNTIF({S}!H2:H200,"FAIL")'],
              ["BLOCKED", f'=COUNTIF({S}!H2:H200,"BLOCKED")'],
              ["N/A", f'=COUNTIF({S}!H2:H200,"N/A")'],
              ["Nieuzupelnione", f'=COUNTBLANK({S}!H2:H200)-COUNTBLANK({S}!A2:A200)']]:
        wp.append(r)
    wp.column_dimensions["A"].width = 20
    wp.column_dimensions["B"].width = 48
    for c in wp[1]:
        c.font = Font(bold=True)
        c.fill = hdr_fill

    wd = wb.create_sheet("Dane testowe")
    for r in DANE:
        wd.append(r)
    wd.column_dimensions["A"].width = 30
    wd.column_dimensions["B"].width = 80
    for row in wd.iter_rows():
        row[0].font = Font(bold=True)
        for c in row:
            c.alignment = Alignment(vertical="top", wrap_text=True)

    out = f"EdSharp_QA_Retest_{WERSJA}.xlsx"
    wb.save(out)
    print("Zapisano:", out, "przypadkow:", len(CASES))


if __name__ == "__main__":
    build()
